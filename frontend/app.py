import streamlit as st
import requests
from fpdf import FPDF
from io import BytesIO
from docx import Document

# ==========================
# Backend URL
# ==========================
BACKEND_URL = "https://legalease-production-48e4.up.railway.app"

st.set_page_config(page_title="LegalEase AI", page_icon="⚖️")

st.title("⚖️ LegalEase AI")
st.write("Generate, Edit and Download Professional Legal Documents using AI.")

# ==========================
# User Inputs
# ==========================
document_type = st.selectbox(
    "Document Type",
    [
        "Rental Agreement",
        "Employment Contract",
        "NDA",
        "Service Agreement"
    ]
)

parties = st.text_area("Parties")

terms = st.text_area("Terms & Conditions")

effective_date = st.date_input("Effective Date")

# ==========================
# Generate Button
# ==========================
if st.button("Generate Document"):

    data = {
        "document_type": document_type,
        "parties": parties,
        "terms": terms,
        "effective_date": str(effective_date)
    }

    try:

        response = requests.post(
            f"{BACKEND_URL}/generate",
            json=data,
            timeout=30
        )

        if response.status_code == 200:

            result = response.json()

            st.success("✅ Document Generated Successfully!")

            # ==========================
            # Editable Document
            # ==========================
            edited_document = st.text_area(
                "✏️ Edit Document",
                value=result["document"],
                height=450
            )

            st.info("You can edit the generated document before downloading.")

            # ==========================
            # TXT Download
            # ==========================
            st.download_button(
                label="📄 Download TXT",
                data=edited_document,
                file_name="LegalEase_Document.txt",
                mime="text/plain"
            )

            # ==========================
            # PDF Download
            # ==========================
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)

            for line in edited_document.split("\n"):
                pdf.multi_cell(0, 10, line)

            pdf_bytes = pdf.output(dest="S").encode("latin-1")

            st.download_button(
                label="📕 Download PDF",
                data=pdf_bytes,
                file_name="LegalEase_Document.pdf",
                mime="application/pdf"
            )

            # ==========================
            # DOCX Download
            # ==========================
            doc = Document()

            doc.add_heading("LegalEase AI", level=1)
            doc.add_paragraph(edited_document)

            buffer = BytesIO()

            doc.save(buffer)

            buffer.seek(0)

            st.download_button(
                label="📘 Download DOCX",
                data=buffer,
                file_name="LegalEase_Document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

        else:
            st.error(response.text)

    except Exception as e:
        st.error(f"❌ Cannot connect to backend:\n\n{e}")